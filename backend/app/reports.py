"""Export Excel reporting for management and offline competition review."""
from __future__ import annotations

import io
from typing import Dict, Any, List
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def build_excel_report(report_data: Dict[str, Any]) -> io.BytesIO:
    """Generate styled workbook summarizing Predict-Detect-Simulate-Recommend outputs."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default sheet

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1B1464", end_color="1B1464", fill_type="solid")
    sub_fill = PatternFill(start_color="2E228F", end_color="2E228F", fill_type="solid")
    bold_font = Font(name="Segoe UI", size=10, bold=True)
    normal_font = Font(name="Segoe UI", size=10)
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    right = Alignment(horizontal="right", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )

    # Sheet 1: Tổng quan điều hành
    ws1 = wb.create_sheet(title="Tổng quan điều hành")
    ws1.views.sheetView[0].showGridLines = True
    ws1.append(["BÁO CÁO ĐIỀU HÀNH LOGISTICS NHÀ MÁY (D2 CONTROL ROOM)"])
    ws1.cell(1, 1).font = Font(name="Segoe UI", size=14, bold=True, color="1B1464")
    ws1.append(["Ngày xuất báo cáo:", pd.Timestamp.now().strftime("%Y-%m-%d %H:%M:%S")])
    ws1.append(["Phiên bản dữ liệu:", report_data.get("data_version", "SYNTHETIC")])
    ws1.append(["Thời điểm ra quyết định (as_of):", report_data.get("as_of", "")])
    ws1.append([])

    ws1.append(["KHUYẾN NGHỊ ĐIỀU PHỐI TỐI ƯU"])
    ws1.cell(6, 1).font = Font(name="Segoe UI", size=12, bold=True, color="FF6B00")
    ws1.append([report_data.get("recommendation_statement", "N/A")])
    ws1.cell(7, 1).font = Font(name="Segoe UI", size=10, italic=True)
    ws1.append([])

    # Sheet 2: So sánh phương án
    ws2 = wb.create_sheet(title="So sánh phương án")
    ws2.views.sheetView[0].showGridLines = True
    actions_headers = [
        "Mã phương án", "Tên phương án", "Chi phí can thiệp (VND)", "Tổn thất kỳ vọng E[L] (VND)",
        "Rủi ro đuôi xấu CVaR90 (VND)", "Điểm tổng hợp rủi ro", "Phút trễ TB (phút)", "Tỷ lệ trễ TB (%)",
        "Khả thi?", "Tập tối ưu Pareto"
    ]
    ws2.append(actions_headers)
    for col_idx in range(1, len(actions_headers) + 1):
        cell = ws2.cell(1, col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    for act in report_data.get("evaluations", []):
        ws2.append([
            act.get("action_id"),
            act.get("name"),
            act.get("cost_vnd", 0),
            act.get("expected_loss_vnd", 0),
            act.get("cvar90_loss_vnd", 0),
            act.get("risk_score", 0),
            act.get("avg_late_minutes", 0),
            act.get("avg_late_rate", 0),
            "Có" if act.get("is_feasible") else f"Không ({act.get('infeasibility_reason')})",
            "Đạt Pareto" if act.get("is_pareto") else "Bị trội",
        ])

    # Sheet 3: Dự báo tải & Năng lực
    ws3 = wb.create_sheet(title="Dự báo tải & Năng lực")
    ws3.views.sheetView[0].showGridLines = True
    pred_headers = [
        "Khung giờ", "Ca", "Dự báo q10 (phút)", "Dự báo q50 (phút)", "Dự báo q90 (phút)",
        "Năng lực khả dụng (phút)", "Picking", "Loading", "Transport", "Tỷ lệ Tải/Năng lực", "Cảnh báo rủi ro"
    ]
    ws3.append(pred_headers)
    for col_idx in range(1, len(pred_headers) + 1):
        cell = ws3.cell(1, col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    for pt in report_data.get("timeline", []):
        ws3.append([
            pt.get("bucket_start"),
            pt.get("shift_id"),
            pt.get("q10"),
            pt.get("q50"),
            pt.get("q90"),
            pt.get("capacity"),
            pt.get("capacity_picking"),
            pt.get("capacity_loading"),
            pt.get("capacity_transport"),
            f"{pt.get('overload_ratio', 0)*100:.0f}%",
            "CẢNH BÁO" if pt.get("is_risk") else "Bình thường",
        ])

    # Sheet 4: Kiểm chứng Backtest
    ws4 = wb.create_sheet(title="Kiểm chứng Backtest")
    ws4.views.sheetView[0].showGridLines = True
    ws4.append(["Chỉ số đánh giá", "Mô hình Baseline (Seasonal Naive)", "LightGBM Quantile (Raw)", "LightGBM (Conformal Calibrated)"])
    for col_idx in range(1, 5):
        cell = ws4.cell(1, col_idx)
        cell.font = header_font
        cell.fill = sub_fill
        cell.alignment = center

    bt = report_data.get("backtest", {})
    b_m = bt.get("baseline", {})
    l_raw = bt.get("lightgbm_raw", {})
    l_cal = bt.get("lightgbm_calibrated", {})

    metrics_rows = [
        ("WAPE (Sai số tỷ đối tuyệt đối)", b_m.get("wape"), l_raw.get("wape"), l_cal.get("wape")),
        ("MAE (Sai số tuyệt đối - phút)", b_m.get("mae"), l_raw.get("mae"), l_cal.get("mae")),
        ("Bias (Độ lệch)", b_m.get("bias"), l_raw.get("bias"), l_cal.get("bias")),
        ("Độ bao phủ thực tế [q10, q90]", f"{b_m.get('coverage', 0)*100:.1f}%", f"{l_raw.get('coverage', 0)*100:.1f}%", f"{l_cal.get('coverage', 0)*100:.1f}%"),
        ("Độ rộng dải dự báo (phút)", b_m.get("interval_width"), l_raw.get("interval_width"), l_cal.get("interval_width")),
        ("Pinball Loss (q50)", b_m.get("pinball_50"), l_raw.get("pinball_50"), l_cal.get("pinball_50")),
    ]
    for m_label, b_val, r_val, c_val in metrics_rows:
        ws4.append([m_label, b_val, r_val, c_val])

    # Auto-adjust column widths for all sheets
    for ws in [ws1, ws2, ws3, ws4]:
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
